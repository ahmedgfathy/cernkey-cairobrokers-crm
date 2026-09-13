import csv
import io
import re
import unicodedata
from datetime import datetime

def export_csv(queryset, fields, header_labels, filename):
    response = io.StringIO()
    writer = csv.writer(response)
    writer.writerow(header_labels)
    for obj in queryset:
        row = []
        for field in fields:
            val = obj
            for part in field.split('__'):
                if val is None:
                    val = ''
                    break
                val = getattr(val, part, '')
            if isinstance(val, bool):
                val = 'Yes' if val else 'No'
            elif isinstance(val, datetime):
                val = val.strftime('%Y-%m-%d %H:%M')
            elif hasattr(val, 'strftime'):
                val = val.strftime('%Y-%m-%d')
            elif hasattr(val, 'pk'):
                val = str(val)
            row.append(val if val is not None else '')
        writer.writerow(row)
    from django.http import HttpResponse
    content = '\ufeff' + response.getvalue()
    http_response = HttpResponse(content.encode('utf-8'), content_type='text/csv; charset=utf-8')
    http_response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    return http_response


def export_excel(queryset, fields, header_labels, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = 'Data'
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2C3E50', end_color='2C3E50', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9'),
    )
    for col_idx, label in enumerate(header_labels, 1):
        cell = ws.cell(row=1, column=col_idx, value=label)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    for row_idx, obj in enumerate(queryset, 2):
        for col_idx, field in enumerate(fields, 1):
            val = obj
            for part in field.split('__'):
                if val is None:
                    val = ''
                    break
                val = getattr(val, part, '')
            if isinstance(val, bool):
                val = 'Yes' if val else 'No'
            elif isinstance(val, datetime):
                val = val.strftime('%Y-%m-%d %H:%M')
            elif hasattr(val, 'strftime'):
                val = val.strftime('%Y-%m-%d')
            elif hasattr(val, 'pk'):
                val = str(val)
            cell = ws.cell(row=row_idx, column=col_idx, value=val if val is not None else '')
            cell.border = thin_border
            cell.alignment = Alignment(vertical='center')
    for col_idx, label in enumerate(header_labels, 1):
        max_length = len(str(label))
        for row_idx in range(2, min(len(queryset) + 2, 52)):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            if cell_val:
                max_length = max(max_length, len(str(cell_val)))
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(max_length + 4, 40)
    from django.http import HttpResponse
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


def parse_uploaded_file(file_obj):
    filename = file_obj.name.lower()
    rows = []
    headers = []
    if filename.endswith('.csv'):
        content = file_obj.read()
        for encoding in ['utf-8-sig', 'utf-8', 'latin-1', 'cp1256']:
            try:
                text = content.decode(encoding)
                break
            except (UnicodeDecodeError, LookupError):
                continue
        reader = csv.reader(io.StringIO(text))
        all_rows = list(reader)
        if not all_rows:
            return [], []
        headers = [h.strip() for h in all_rows[0]]
        rows = all_rows[1:]
    elif filename.endswith('.xlsx'):
        from openpyxl import load_workbook

        wb = load_workbook(file_obj, read_only=True, data_only=True)
        ws = wb.active
        for row_idx, row in enumerate(ws.iter_rows(values_only=True)):
            if row_idx == 0:
                headers = [str(h).strip() if h else '' for h in row]
            else:
                rows.append([str(c) if c is not None else '' for c in row])
        wb.close()
    return headers, rows


def auto_match_headers(file_headers, db_fields):
    def normalize(value):
        value = unicodedata.normalize('NFKC', str(value)).casefold()
        value = re.sub(r'[^\w]+', ' ', value, flags=re.UNICODE)
        return ' '.join(value.split())

    mapping = {}
    normalized_db = {}
    for key, label in db_fields.items():
        normalized_db[normalize(key)] = key
        normalized_db[normalize(label)] = key
    for idx, header in enumerate(file_headers):
        h = normalize(header)
        if h in normalized_db:
            mapping[idx] = normalized_db[h]
        else:
            for norm_key, field_key in normalized_db.items():
                if h and (h in norm_key or norm_key in h):
                    mapping[idx] = field_key
                    break
    return mapping
